"""
RAG System with Multiple Agents & Conversation Memory
=======================================================
- Create multiple agents with dedicated data and memory storage
- Conversation history for each agent
- Agent details stored in Excel file
- Embeddings powered by custom trained EmbeddingModel
"""

import os
import sys
import json
import time
import uuid
import asyncio
import logging
import httpx
import shutil
from pathlib import Path
from typing import List, Optional, Dict
from datetime import datetime

import torch
import torch.nn as nn
import torch.nn.functional as F
import numpy as np
import faiss
from openpyxl import Workbook, load_workbook

from fastapi import FastAPI, UploadFile, File, HTTPException, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field

from langchain_community.document_loaders import (
    CSVLoader,
    PyPDFLoader,
    TextLoader,
    UnstructuredHTMLLoader,
    UnstructuredWordDocumentLoader,
)
from langchain_text_splitters import RecursiveCharacterTextSplitter

# ─────────────────────────────────────────────
#  Add project root to path so model/ and tokenizer.py are importable
# ─────────────────────────────────────────────
PROJECT_ROOT = os.path.abspath(os.path.dirname(__file__))
sys.path.insert(0, PROJECT_ROOT)

from tokenizer import MyTokenizer
from model.full_model import FullModel

# ─────────────────────────────────────────────
#  EmbeddingModel (same as training)
# ─────────────────────────────────────────────

class EmbeddingModel(nn.Module):

    def __init__(self, vocab_size=10000):
        super().__init__()
        self.base_model = FullModel(vocab_size=vocab_size)

    def forward(self, input_ids):
        x = self.base_model.embeddings(input_ids)
        x = self.base_model.transformer(x)

        # Masked mean pooling
        mask = (input_ids != 0).unsqueeze(-1).float()
        x = x * mask
        x = x.sum(dim=1) / mask.sum(dim=1).clamp(min=1e-9)

        x = F.normalize(x, p=2, dim=1)
        return x

# ─────────────────────────────────────────────
#  Configuration
# ─────────────────────────────────────────────
OLLAMA_BASE_URL  = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434").rstrip("/")
OLLAMA_MODEL     = os.getenv("OLLAMA_MODEL", "llama3.1")
OLLAMA_TIMEOUT   = float(os.getenv("OLLAMA_TIMEOUT_SECONDS", "120"))
OLLAMA_ENABLED   = os.getenv("OLLAMA_ENABLED", "true").lower() in {"1", "true", "yes", "on"}

MODEL_PATH       = os.getenv("EMBED_MODEL_PATH", os.path.join(PROJECT_ROOT, "embedding_model_best.pt"))
VOCAB_SIZE       = int(os.getenv("VOCAB_SIZE", "10000"))
EMBED_DIM        = 384                         # must match your trained model

CHUNK_SIZE       = int(os.getenv("CHUNK_SIZE", "500"))
CHUNK_OVERLAP    = int(os.getenv("CHUNK_OVERLAP", "50"))
TOP_K            = int(os.getenv("TOP_K", "5"))
BATCH_SIZE       = int(os.getenv("EMBED_BATCH_SIZE", "32"))

BASE_DIR         = Path(os.getenv("RAG_BASE_DIR", "rag_data"))
AGENTS_EXCEL     = BASE_DIR / "agents.xlsx"

SUPPORTED_EXTENSIONS = {".csv", ".pdf", ".txt", ".md", ".html", ".htm", ".docx"}

# ─────────────────────────────────────────────
#  Logging
# ─────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("rag")

# ─────────────────────────────────────────────
#  Globals
# ─────────────────────────────────────────────
embed_model: Optional[EmbeddingModel] = None
tokenizer: Optional[MyTokenizer]      = None
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

agent_cache: Dict[str, tuple] = {}

# ─────────────────────────────────────────────
#  Embedding helper
# ─────────────────────────────────────────────

def _pad_batch(id_list: List[List[int]]) -> torch.Tensor:
    """Pad a list of token-id lists to the same length and return a LongTensor."""
    max_len = max(len(x) for x in id_list)
    padded  = []
    for x in id_list:
        pad_len = max_len - len(x)
        padded.append(x + [0] * pad_len)
    return torch.tensor(padded, dtype=torch.long)


def encode_texts(texts: List[str]) -> np.ndarray:
    """
    Encode a list of strings → float32 numpy array of shape (N, 384).
    Processes in batches so large document sets don't OOM.
    """
    all_vecs = []

    for start in range(0, len(texts), BATCH_SIZE):
        batch_texts = texts[start : start + BATCH_SIZE]

        # Tokenise
        id_list = [tokenizer.encode(t) for t in batch_texts]

        # Safety: if any sequence is empty give it a single PAD token
        id_list = [ids if ids else [0] for ids in id_list]

        input_ids = _pad_batch(id_list).to(device)

        with torch.no_grad():
            vecs = embed_model(input_ids)          # (B, 384), L2-normalised

        all_vecs.append(vecs.cpu().numpy())

    return np.concatenate(all_vecs, axis=0).astype("float32")

# ─────────────────────────────────────────────
#  FastAPI App
# ─────────────────────────────────────────────
app = FastAPI(
    title="RAG System with Multiple Agents",
    description="Retrieval-Augmented Generation with Ollama, FAISS, and custom EmbeddingModel",
    version="2.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─────────────────────────────────────────────
#  System init helpers
# ─────────────────────────────────────────────
def initialize_system():
    BASE_DIR.mkdir(parents=True, exist_ok=True)
    if not AGENTS_EXCEL.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Agents"
        ws.append(["Agent ID", "Agent Name", "Created At", "Description"])
        wb.save(AGENTS_EXCEL)


def get_agent_dirs(agent_id: str) -> Dict[str, Path]:
    agent_dir = BASE_DIR / agent_id
    return {
        "agent":  agent_dir,
        "data":   agent_dir / "data",
        "index":  agent_dir / "index",
        "memory": agent_dir / "memory",
    }


def create_agent_folders(agent_id: str):
    for d in get_agent_dirs(agent_id).values():
        d.mkdir(parents=True, exist_ok=True)


def get_agents_from_excel() -> List[Dict]:
    if not AGENTS_EXCEL.exists():
        return []
    wb = load_workbook(AGENTS_EXCEL)
    ws = wb["Agents"]
    agents = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            agents.append({
                "agent_id":   row[0],
                "agent_name": row[1],
                "created_at": row[2],
                "description": row[3] if len(row) > 3 else "",
            })
    wb.close()
    return agents


def add_agent_to_excel(agent_id: str, agent_name: str, description: str = ""):
    wb = load_workbook(AGENTS_EXCEL)
    ws = wb["Agents"]
    ws.append([agent_id, agent_name, datetime.now().isoformat(), description])
    wb.save(AGENTS_EXCEL)
    wb.close()


def delete_agent_from_excel(agent_id: str):
    if not AGENTS_EXCEL.exists():
        return
    wb = load_workbook(AGENTS_EXCEL)
    ws = wb["Agents"]
    rows_to_delete = []
    for idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        if row[0].value == agent_id:
            rows_to_delete.append(idx)
    for idx in reversed(rows_to_delete):
        ws.delete_rows(idx)
    wb.save(AGENTS_EXCEL)
    wb.close()

# ─────────────────────────────────────────────
#  Index load / save
# ─────────────────────────────────────────────
def _load_index_from_disk(agent_id: str):
    dirs        = get_agent_dirs(agent_id)
    index_file  = dirs["index"] / "faiss.index"
    meta_file   = dirs["index"] / "metadata.json"
    state_file  = dirs["index"] / "state.json"

    faiss_index  = None
    chunk_store  = []
    index_status = {"state": "empty", "total_chunks": 0, "files": []}

    if index_file.exists() and meta_file.exists():
        try:
            faiss_index = faiss.read_index(str(index_file))
            with open(meta_file, "r", encoding="utf-8") as f:
                chunk_store = json.load(f)
            if state_file.exists():
                with open(state_file, "r") as f:
                    index_status = json.load(f)
            log.info(f"Loaded index for agent {agent_id}: {len(chunk_store)} chunks")
        except Exception as e:
            log.warning(f"Could not load index for agent {agent_id}: {e}")

    agent_cache[agent_id] = (faiss_index, chunk_store, index_status)
    return faiss_index, chunk_store, index_status


def _save_index_to_disk(agent_id: str, faiss_index, chunk_store, index_status):
    dirs = get_agent_dirs(agent_id)
    faiss.write_index(faiss_index, str(dirs["index"] / "faiss.index"))
    with open(dirs["index"] / "metadata.json", "w", encoding="utf-8") as f:
        json.dump(chunk_store, f, ensure_ascii=False, indent=2)
    with open(dirs["index"] / "state.json", "w") as f:
        json.dump(index_status, f, indent=2)

# ─────────────────────────────────────────────
#  Conversation memory
# ─────────────────────────────────────────────
def save_conversation_memory(agent_id: str, conversation_id: str, messages: List[Dict]):
    dirs = get_agent_dirs(agent_id)
    mem_file = dirs["memory"] / f"{conversation_id}.json"
    with open(mem_file, "w", encoding="utf-8") as f:
        json.dump({
            "conversation_id": conversation_id,
            "agent_id":        agent_id,
            "messages":        messages,
            "updated_at":      datetime.now().isoformat(),
        }, f, ensure_ascii=False, indent=2)


def load_conversation_memory(agent_id: str, conversation_id: str) -> List[Dict]:
    dirs     = get_agent_dirs(agent_id)
    mem_file = dirs["memory"] / f"{conversation_id}.json"
    if mem_file.exists():
        with open(mem_file, "r", encoding="utf-8") as f:
            return json.load(f).get("messages", [])
    return []


def list_conversations(agent_id: str) -> List[Dict]:
    dirs          = get_agent_dirs(agent_id)
    conversations = []
    if dirs["memory"].exists():
        for f in dirs["memory"].glob("*.json"):
            with open(f, "r", encoding="utf-8") as fp:
                data = json.load(fp)
                conversations.append({
                    "conversation_id": data.get("conversation_id"),
                    "updated_at":      data.get("updated_at"),
                    "message_count":   len(data.get("messages", [])),
                })
    conversations.sort(key=lambda x: x["updated_at"], reverse=True)
    return conversations

# ─────────────────────────────────────────────
#  Document loader
# ─────────────────────────────────────────────
def _load_file(path: Path) -> list:
    ext = path.suffix.lower()
    try:
        if ext == ".csv":
            loader = CSVLoader(str(path), encoding="utf-8")
        elif ext == ".pdf":
            loader = PyPDFLoader(str(path))
        elif ext in {".html", ".htm"}:
            loader = UnstructuredHTMLLoader(str(path))
        elif ext == ".docx":
            loader = UnstructuredWordDocumentLoader(str(path))
        else:
            loader = TextLoader(str(path), encoding="utf-8")
        return loader.load()
    except Exception as e:
        log.error(f"Failed to load {path.name}: {e}")
        return []

# ─────────────────────────────────────────────
#  Response format detection
# ─────────────────────────────────────────────
def detect_response_format(question: str) -> str:
    q = question.lower()
    format_keywords = {
        "table":       ["table", "compare", "comparison", "vs", "versus", "difference between"],
        "bullet_list": ["list", "features", "points", "bullet", "what are the"],
        "step_by_step":["step", "how to", "guide", "tutorial", "walkthrough"],
        "code":        ["code", "example code", "write code", "function"],
        "json":        ["json", "api response"],
        "normal":      ["explain", "what is", "tell me about"],
    }
    for fmt, keywords in format_keywords.items():
        if any(kw in q for kw in keywords):
            return fmt
    return "normal"

# ─────────────────────────────────────────────
#  Pydantic models
# ─────────────────────────────────────────────
class CreateAgentRequest(BaseModel):
    agent_name:  str           = Field(..., description="Name of the agent")
    description: Optional[str] = Field("",  description="Description of the agent")


class QueryRequest(BaseModel):
    question:        str
    conversation_id: Optional[str] = Field(None, description="ID of the conversation to continue")
    top_k:           int   = TOP_K
    temperature:     float = 0.0


class QueryResponse(BaseModel):
    answer:          str
    format:          str
    conversation_id: str
    sources:         List[dict]
    model:           str
    latency_ms:      float


class IndexResponse(BaseModel):
    message:       str
    total_chunks:  int
    files_indexed: List[str]

# ─────────────────────────────────────────────
#  Startup — load custom embedding model
# ─────────────────────────────────────────────
@app.on_event("startup")
async def startup():
    global embed_model, tokenizer

    initialize_system()

    # ── Load tokenizer ──────────────────────────────
    log.info("Loading tokenizer...")
    tokenizer = MyTokenizer()
    log.info("Tokenizer loaded ✓")

    # ── Load EmbeddingModel ─────────────────────────
    log.info("Loading EmbeddingModel from: %s", MODEL_PATH)

    if not os.path.exists(MODEL_PATH):
        raise RuntimeError(
            f"Model file not found: {MODEL_PATH}\n"
            f"Copy embedding_model_best.pt into {PROJECT_ROOT}"
        )

    embed_model = EmbeddingModel(vocab_size=VOCAB_SIZE)

    state_dict = torch.load(
        MODEL_PATH,
        map_location=device,
        weights_only=True
    )
    embed_model.load_state_dict(state_dict)
    embed_model.to(device)
    embed_model.eval()

    log.info("EmbeddingModel loaded on %s ✓", device)

# ─────────────────────────────────────────────
#  Agent management endpoints
# ─────────────────────────────────────────────
@app.post("/agents", summary="Create a new agent")
async def create_agent(req: CreateAgentRequest):
    agent_id = str(uuid.uuid4())
    create_agent_folders(agent_id)
    add_agent_to_excel(agent_id, req.agent_name, req.description)
    _load_index_from_disk(agent_id)
    return {
        "agent_id":   agent_id,
        "agent_name": req.agent_name,
        "description":req.description,
        "message":    "Agent created successfully",
    }


@app.get("/agents", summary="List all agents")
async def list_agents():
    return {"agents": get_agents_from_excel()}


@app.delete("/agents/{agent_id}", summary="Delete an agent")
async def delete_agent(agent_id: str):
    dirs = get_agent_dirs(agent_id)
    if dirs["agent"].exists():
        try:
            shutil.rmtree(dirs["agent"])
        except Exception as e:
            log.warning(f"Could not fully delete agent dir {agent_id}: {e}")
    delete_agent_from_excel(agent_id)
    agent_cache.pop(agent_id, None)
    return {"message": f"Agent {agent_id} deleted successfully"}

# ─────────────────────────────────────────────
#  File management endpoints
# ─────────────────────────────────────────────
@app.post("/agents/{agent_id}/upload", summary="Upload files to an agent")
async def upload_files(agent_id: str, files: List[UploadFile] = File(...)):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")

    saved, skipped = [], []
    for file in files:
        ext = Path(file.filename).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            skipped.append({"filename": file.filename, "reason": f"Unsupported extension '{ext}'"})
            continue
        dest = dirs["data"] / file.filename
        with open(dest, "wb") as f:
            shutil.copyfileobj(file.file, f)
        saved.append(file.filename)

    return {
        "uploaded":              saved,
        "skipped":               skipped,
        "total_files_in_store":  len(list(dirs["data"].iterdir())),
        "hint":                  "Call POST /agents/{agent_id}/index to rebuild the index",
    }


@app.get("/agents/{agent_id}/files", summary="List uploaded files for an agent")
async def list_files(agent_id: str):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")

    files = []
    if dirs["data"].exists():
        for p in sorted(dirs["data"].iterdir()):
            files.append({
                "filename":  p.name,
                "size_kb":   round(p.stat().st_size / 1024, 2),
                "extension": p.suffix.lower(),
            })
    return {"files": files, "count": len(files)}


@app.delete("/agents/{agent_id}/files/{filename}", summary="Delete a file from an agent")
async def delete_file(agent_id: str, filename: str):
    dirs   = get_agent_dirs(agent_id)
    target = dirs["data"] / filename
    if not target.exists():
        raise HTTPException(status_code=404, detail="File not found")
    target.unlink()
    return {"message": f"'{filename}' deleted. Call /index to rebuild."}

# ─────────────────────────────────────────────
#  Index endpoint
# ─────────────────────────────────────────────
@app.post("/agents/{agent_id}/index", response_model=IndexResponse, summary="Build index for an agent")
async def build_index(agent_id: str):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")
    if not dirs["data"].exists() or not list(dirs["data"].iterdir()):
        raise HTTPException(status_code=400, detail="No files uploaded yet")

    result = await asyncio.get_event_loop().run_in_executor(None, _build_index_sync, agent_id)
    return result


def _build_index_sync(agent_id: str) -> IndexResponse:
    dirs = get_agent_dirs(agent_id)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ",", " ", ""],
    )

    all_chunks:    List[dict] = []
    files_indexed: List[str]  = []

    for file_path in sorted(dirs["data"].iterdir()):
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        log.info(f"Indexing {file_path.name} for agent {agent_id}")
        docs = _load_file(file_path)
        if not docs:
            continue
        chunks = splitter.split_documents(docs)
        for chunk in chunks:
            all_chunks.append({
                "chunk_id": str(uuid.uuid4()),
                "text":     chunk.page_content.strip(),
                "source":   file_path.name,
                "page":     chunk.metadata.get("page", 0),
                "row":      chunk.metadata.get("row", None),
            })
        files_indexed.append(file_path.name)
        log.info(f"  → {len(chunks)} chunks from {file_path.name}")

    if not all_chunks:
        raise HTTPException(status_code=500, detail="No text could be extracted")

    # ── Embed with custom model ──────────────────────
    log.info(f"Embedding {len(all_chunks)} chunks for agent {agent_id}...")
    texts   = [c["text"] for c in all_chunks]
    vectors = encode_texts(texts)             # (N, 384) float32, L2-normalised

    # ── Build FAISS index ────────────────────────────
    faiss_index = faiss.IndexFlatIP(EMBED_DIM)   # inner product == cosine sim when normalised
    faiss_index.add(vectors)

    index_status = {
        "state":        "ready",
        "total_chunks": len(all_chunks),
        "files":        files_indexed,
        "embed_model":  "custom EmbeddingModel (embedding_model_best.pt)",
        "chunk_size":   CHUNK_SIZE,
        "chunk_overlap":CHUNK_OVERLAP,
        "built_at":     datetime.now().isoformat(),
    }

    _save_index_to_disk(agent_id, faiss_index, all_chunks, index_status)
    agent_cache[agent_id] = (faiss_index, all_chunks, index_status)

    log.info(f"Index built: {len(all_chunks)} chunks, {len(files_indexed)} files")
    return IndexResponse(
        message="Index built successfully",
        total_chunks=len(all_chunks),
        files_indexed=files_indexed,
    )


@app.delete("/agents/{agent_id}/index", summary="Delete index for an agent")
async def delete_index(agent_id: str):
    dirs = get_agent_dirs(agent_id)
    agent_cache[agent_id] = (None, [], {"state": "empty", "total_chunks": 0, "files": []})
    for fname in ["faiss.index", "metadata.json", "state.json"]:
        (dirs["index"] / fname).unlink(missing_ok=True)
    return {"message": "Index deleted. Upload files and call /index to rebuild."}


@app.get("/agents/{agent_id}/status", summary="Get agent index status")
async def get_status(agent_id: str):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")
    if agent_id not in agent_cache:
        _load_index_from_disk(agent_id)
    _, _, index_status = agent_cache[agent_id]
    return {
        **index_status,
        "ollama_url":   OLLAMA_BASE_URL,
        "ollama_model": OLLAMA_MODEL,
        "embed_model":  "custom EmbeddingModel (embedding_model_best.pt)",
        "device":       str(device),
        "data_dir":     str(dirs["data"]),
        "index_dir":    str(dirs["index"]),
    }

# ─────────────────────────────────────────────
#  Query endpoint
# ─────────────────────────────────────────────
@app.post("/agents/{agent_id}/query", response_model=QueryResponse, summary="Query an agent")
async def query(agent_id: str, req: QueryRequest):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")

    if agent_id not in agent_cache:
        _load_index_from_disk(agent_id)
    faiss_index, chunk_store, _ = agent_cache[agent_id]

    if faiss_index is None or not chunk_store:
        raise HTTPException(status_code=400, detail="Index is empty. Upload files and call /index first.")

    t0 = time.perf_counter()

    # ── Load / create conversation ───────────────────
    conversation_id = req.conversation_id or str(uuid.uuid4())
    messages        = load_conversation_memory(agent_id, conversation_id)

    # ── Embed query ──────────────────────────────────
    q_vec = encode_texts([req.question])          # (1, 384)

    # ── FAISS search ─────────────────────────────────
    k = min(req.top_k, len(chunk_store))
    scores, indices = faiss_index.search(q_vec, k)

    retrieved = []
    for score, idx in zip(scores[0], indices[0]):
        if idx < 0:
            continue
        chunk = chunk_store[idx]
        retrieved.append({
            "text":   chunk["text"],
            "source": chunk["source"],
            "page":   chunk["page"],
            "score":  float(score),
        })

    if not retrieved:
        raise HTTPException(status_code=500, detail="No relevant chunks found")

    # ── Build prompt ─────────────────────────────────
    detected_format   = detect_response_format(req.question)
    context_blocks    = "\n\n".join(
        [f"[Source: {r['source']} | Score: {r['score']:.3f}]\n{r['text']}" for r in retrieved]
    )
    history_str = "\n".join(
        [f"{m['role']}: {m['content']}" for m in messages[-10:]]
    )

    format_instructions = {
        "table":       "\nPlease format your response as a Markdown table.",
        "bullet_list": "\nPlease format your response as a Markdown bullet list.",
        "step_by_step":"\nPlease format your response as a numbered step-by-step guide.",
        "code":        "\nPlease include code examples with proper Markdown code blocks.",
    }.get(detected_format, "")

    prompt = f"""You are a helpful assistant. Use the conversation history and provided context to answer the question.{format_instructions}

Conversation History:
{history_str}

Context:
{context_blocks}

Question: {req.question}
Answer:"""

    # ── Print context and prompt before LLM ──────────
    print("\n" + "="*80)
    print("CONTEXT:")
    print(context_blocks)
    print("\n" + "="*80)
    print("PROMPT:")
    print(prompt)
    print("="*80 + "\n")

    # ── Call Ollama ──────────────────────────────────
    if not OLLAMA_ENABLED:
        answer = "[Ollama disabled] Context retrieved but LLM call skipped."
    else:
        answer = await _call_ollama(prompt, req.temperature)

    # ── Save conversation ────────────────────────────
    messages.append({"role": "user",      "content": req.question})
    messages.append({"role": "assistant", "content": answer})
    save_conversation_memory(agent_id, conversation_id, messages)

    latency_ms = (time.perf_counter() - t0) * 1000

    return QueryResponse(
        answer=answer,
        format=detected_format,
        conversation_id=conversation_id,
        sources=retrieved,
        model=OLLAMA_MODEL,
        latency_ms=round(latency_ms, 1),
    )

# ─────────────────────────────────────────────
#  Conversation endpoints
# ─────────────────────────────────────────────
@app.get("/agents/{agent_id}/conversations", summary="List conversations for an agent")
async def get_conversations(agent_id: str):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")
    return {"conversations": list_conversations(agent_id)}


@app.get("/agents/{agent_id}/conversations/{conversation_id}", summary="Get a specific conversation")
async def get_conversation(agent_id: str, conversation_id: str):
    dirs = get_agent_dirs(agent_id)
    if not dirs["agent"].exists():
        raise HTTPException(status_code=404, detail="Agent not found")
    messages = load_conversation_memory(agent_id, conversation_id)
    return {"agent_id": agent_id, "conversation_id": conversation_id, "messages": messages}


@app.delete("/agents/{agent_id}/conversations/{conversation_id}", summary="Delete a conversation")
async def delete_conversation(agent_id: str, conversation_id: str):
    dirs     = get_agent_dirs(agent_id)
    mem_file = dirs["memory"] / f"{conversation_id}.json"
    if mem_file.exists():
        mem_file.unlink()
    return {"message": "Conversation deleted"}

# ─────────────────────────────────────────────
#  Ollama call
# ─────────────────────────────────────────────
async def _call_ollama(prompt: str, temperature: float) -> str:
    payload = {
        "model":  OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "options": {
            "temperature": temperature,
            "num_predict": 1024,
        },
    }
    try:
        async with httpx.AsyncClient(timeout=OLLAMA_TIMEOUT) as client:
            resp = await client.post(f"{OLLAMA_BASE_URL}/api/generate", json=payload)
            resp.raise_for_status()
            return resp.json().get("response", "").strip()
    except httpx.ConnectError:
        raise HTTPException(status_code=503, detail=f"Cannot connect to Ollama at {OLLAMA_BASE_URL}")
    except httpx.TimeoutException:
        raise HTTPException(status_code=504, detail="Ollama timeout")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ollama error: {e}")

# ─────────────────────────────────────────────
#  Root & health
# ─────────────────────────────────────────────
@app.get("/", include_in_schema=False)
async def root():
    return JSONResponse({
        "message": "RAG System API (Multi-Agent) is running",
        "docs":    "/docs",
        "endpoints": {
            "agents":        "/agents",
            "upload":        "/agents/{agent_id}/upload",
            "index":         "/agents/{agent_id}/index",
            "query":         "/agents/{agent_id}/query",
            "conversations": "/agents/{agent_id}/conversations",
        },
    })


@app.get("/health", summary="Health check")
async def health():
    ollama_ok = False
    try:
        async with httpx.AsyncClient(timeout=3) as client:
            r = await client.get(f"{OLLAMA_BASE_URL}/api/tags")
            ollama_ok = r.status_code == 200
    except Exception:
        pass
    return {
        "status":             "ok",
        "embed_model_loaded": embed_model is not None,
        "embed_model":        "custom EmbeddingModel (embedding_model_best.pt)",
        "device":             str(device),
        "ollama_reachable":   ollama_ok,
        "ollama_url":         OLLAMA_BASE_URL,
        "ollama_model":       OLLAMA_MODEL,
    }

# ─────────────────────────────────────────────
#  Entry point
# ─────────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    uvicorn.run("rag_system:app", host="0.0.0.0", port=8000, reload=True)